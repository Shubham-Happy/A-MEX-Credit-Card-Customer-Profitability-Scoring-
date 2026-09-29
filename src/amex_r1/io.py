from __future__ import annotations

from pathlib import Path
from typing import Mapping

import pandas as pd
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Alignment, Font, PatternFill


FRAMEWORK_FIELDS = [
    "Variables Used",
    "Profitability Equation",
    "Prediction Logic",
    "Variable Selection Logic",
    "Coefficient/Weight Derivation",
    "Feature Transformations",
    "Business Logic",
    "Assumptions",
    "Validation Approach",
    "Additional Notes",
]


def load_dataset(path: str | Path) -> pd.DataFrame:
    path = Path(path)
    return pd.read_csv(path, encoding="utf-8-sig")


def load_predictions_xlsx(path: str | Path) -> pd.DataFrame:
    return pd.read_excel(path, sheet_name="Predictions")


def write_submission_xlsx(
    predictions: pd.DataFrame,
    template_path: str | Path,
    output_path: str | Path,
    framework: Mapping[str, str],
) -> None:
    """Write predictions and framework text into the official-style workbook."""
    template_path = Path(template_path)
    output_path = Path(output_path)

    if template_path.exists():
        wb = load_workbook(template_path)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Predictions"
        wb.create_sheet("Profitability Framework")

    # Ensure required sheets exist.
    if "Predictions" not in wb.sheetnames:
        wb.create_sheet("Predictions", 0)
    if "Profitability Framework" not in wb.sheetnames:
        wb.create_sheet("Profitability Framework")

    ws = wb["Predictions"]
    ws.delete_rows(1, ws.max_row)
    ws.append(["ID", "Prediction"])
    for row in predictions[["ID", "Prediction"]].itertuples(index=False):
        ws.append([int(row.ID), float(row.Prediction)])

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 18
    ws.freeze_panes = "A2"

    fw = wb["Profitability Framework"]
    fw.delete_rows(1, fw.max_row)
    fw.append(["Section", "Response"])
    for field in FRAMEWORK_FIELDS:
        fw.append([field, framework.get(field, "")])

    for cell in fw[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    fw.column_dimensions["A"].width = 34
    fw.column_dimensions["B"].width = 120
    for row in fw.iter_rows(min_row=2, max_col=2):
        row[0].font = Font(bold=True)
        row[0].alignment = Alignment(vertical="top", wrap_text=True)
        row[1].alignment = Alignment(vertical="top", wrap_text=True)
    for idx in range(2, fw.max_row + 1):
        fw.row_dimensions[idx].height = 78

    wb.save(output_path)
